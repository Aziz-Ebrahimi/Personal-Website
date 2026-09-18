from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = "Professor_Website_Information_Checklist.docx"

BLUE = "2E6171"
DARK = "173B45"
MUTED = "66767B"
PALE = "EAF1F2"
BORDER = "CAD7DA"
GREEN = "517A58"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[i]))
            tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)
        borders.append(element)


def set_font(run, size=11, bold=False, color="25383D", italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, 9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_checklist_table(doc, items):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for idx, item in enumerate(items):
        cells = table.add_row().cells
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cells[0], top=110, bottom=110)
        set_cell_margins(cells[1], top=110, bottom=110)
        if idx % 2 == 1:
            set_cell_shading(cells[0], "F7F9F9")
            set_cell_shading(cells[1], "F7F9F9")
        p0 = cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run("☐")
        set_font(r0, 13, color=BLUE)
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(item)
        set_font(r1, 10.5)
    set_table_geometry(table, [520, 8840])
    set_table_borders(table)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_section(doc, number, title, status, intro, items, page_break=False):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{number}. {title}")
    set_font(r, 16, bold=True, color=DARK)
    badge = p.add_run(f"   {status.upper()}")
    set_font(badge, 8.5, bold=True, color=GREEN if status == "Required" else MUTED)
    if intro:
        ip = doc.add_paragraph()
        ip.paragraph_format.keep_with_next = True
        ip.paragraph_format.space_after = Pt(6)
        ir = ip.add_run(intro)
        set_font(ir, 10.5, italic=True, color=MUTED)
    add_checklist_table(doc, items)


sections = [
    ("Basic Profile", "Required", "Core information for the homepage and contact area.", [
        "Full name, including preferred title and middle initial",
        "Academic title (for example, Professor of Plant Genetics)",
        "Department, university, or research center",
        "Preferred professional email address",
        "Office location and mailing address",
        "Professional headshot — JPG or PNG, preferably at least 800 × 800 pixels",
        "Short biography — approximately 150–300 words",
        "Three to six research-interest keywords",
        "Preferred website title",
    ]),
    ("Professional Links", "Optional", "Share only the profiles that should appear publicly.", [
        "Google Scholar profile", "ORCID profile", "University faculty profile",
        "ResearchGate profile", "LinkedIn profile", "GitHub profile",
        "Lab or department website", "Other professional profile links",
    ]),
    ("Research Overview", "Required", "The website currently uses three main research categories.", [
        "Overall research summary — approximately 100–250 words",
        "Confirm or revise category: Remote Sensing",
        "Confirm or revise category: Plant Genetics",
        "Confirm or revise category: AI for Crop Science",
        "Preferred wording, order, or additional categories",
    ]),
    ("Research Projects", "Required", "Provide the following for each project that should appear on the Research page.", [
        "Project title and research category",
        "Short project summary — approximately 100–200 words",
        "Project goals or major research questions",
        "Methods, study region, crops, or datasets involved",
        "Main findings or expected impact",
        "Collaborators and partner organizations",
        "Funding source and grant number, if it should be displayed",
        "Related publication links",
        "One project image, figure, or field photograph — preferably at least 1,200 pixels wide",
        "Image caption, image credit, and confirmation that the image may be published publicly",
    ]),
    ("Publications", "Required", "Preferred: a complete BibTeX (.bib) file exported from a reference manager.", [
        "Complete BibTeX (.bib) file, if available",
        "For each item: title, complete author list, year, and journal or conference",
        "Volume, issue, page numbers, DOI, and public URL",
        "Abstract",
        "Category: Remote Sensing, Genetics, or AI",
        "Mark publications that should appear as Selected Publications on the homepage",
        "Optional PDF, code, dataset, project link, or preview image",
        "Confirm author spelling, author order, dates, and DOI links",
    ]),
    ("Team", "Required", "Provide information for each current member and confirm permission to publish it.", [
        "Full name and role or position",
        "Degree program and expected graduation year, if applicable",
        "Short biography — approximately 50–120 words",
        "Research interests",
        "Professional headshot",
        "Email address or profile link, only if the member wants it displayed",
        "Preferred display order",
        "Confirmation that the photograph and biography may be published publicly",
        "Indicate whether to include alumni, visitors, undergraduates, collaborators, or open positions",
    ]),
    ("Teaching", "Required", "Confirm the current course information and add other courses only if desired.", [
        "Course title: Intro to Plant Genetics",
        "Semester: Spring 2027",
        "Course number and intended student level",
        "Short course description",
        "Syllabus or syllabus link",
        "Office hours",
        "Public course website or learning-platform link",
    ]),
    ("News", "Optional", "For each item, provide the month, year, short announcement, and optional link.", [
        "New publications or grants", "Awards and student achievements",
        "Fieldwork or conference presentations", "Lab announcements",
    ]),
    ("Awards and Honors", "Optional", "Provide one entry for each award that should be displayed.", [
        "Award name", "Awarding organization", "Year", "Optional one-sentence description and link",
    ]),
    ("Professional Service", "Optional", "Examples include editor roles, panels, committees, and university service.", [
        "Role title", "Journal, conference, committee, society, or organization",
        "Dates or years", "Optional short description",
    ]),
    ("CV", "Required", "Please send the latest CV as a Word document or PDF.", [
        "Education and academic appointments", "Research experience", "Publications",
        "Grants and funding", "Teaching", "Awards and professional service",
        "Students advised and professional memberships",
        "Confirm whether the complete CV PDF may be offered as a public download",
    ]),
    ("Website Ownership and Publishing", "Required", "Complete this section before the website is launched.", [
        "Preferred GitHub username or lab/department GitHub organization",
        "Preferred website address",
        "Existing university domain or custom domain, if applicable",
        "Person or organization that should own the website repository",
        "Names or GitHub usernames of people who should have editing access",
        "Confirmation that all public content has been reviewed and approved",
    ]),
    ("Final Review", "Required", "Please verify every item before approving public publication.", [
        "Name, title, department, email, and address are correct",
        "Biography is approved",
        "Publication details are accurate",
        "Team members approved their biographies and photographs",
        "Project images have publication permission and proper credits",
        "Private, unpublished, or sensitive research information has been removed",
        "Funding acknowledgments are correct",
        "All external links work",
        "The website may be made publicly accessible",
    ]),
]

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.78)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for style_name, size, color, before, after in [
    ("Heading 1", 16, DARK, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK, 10, 5),
]:
    st = styles[style_name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

header = section.header
hp = header.paragraphs[0]
hp.text = "PROFESSOR WEBSITE MATERIALS"
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(hp.runs[0], 8.5, bold=True, color=MUTED)
footer = section.footer
add_page_number(footer.paragraphs[0])

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(44)
title.paragraph_format.space_after = Pt(8)
tr = title.add_run("Website Information Checklist")
set_font(tr, 28, bold=True, color=DARK)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(18)
sr = subtitle.add_run("Materials needed to prepare and publish an academic website")
set_font(sr, 13, color=BLUE)

lead = doc.add_table(rows=1, cols=1)
lead.alignment = WD_TABLE_ALIGNMENT.LEFT
lead.autofit = False
cell = lead.cell(0, 0)
set_cell_shading(cell, PALE)
set_cell_margins(cell, top=180, bottom=180, start=220, end=220)
lp = cell.paragraphs[0]
lp.paragraph_format.space_after = Pt(0)
lr = lp.add_run("How to use this document\n")
set_font(lr, 11, bold=True, color=DARK)
lr2 = lp.add_run("Check each completed item and return the document with the requested files. Required items are needed for the first public version; optional items may be added later.")
set_font(lr2, 10.5)
set_table_geometry(lead, [9360])
set_table_borders(lead)

doc.add_paragraph()
for i, (name, status, intro, items) in enumerate(sections, 1):
    page_break = i in {2, 4, 5, 6, 8, 11, 13}
    add_section(doc, i, name, status, intro, items, page_break=page_break)

doc.add_page_break()
closing = doc.add_paragraph(style="Heading 1")
closing.add_run("Suggested Folder Structure")
cp = doc.add_paragraph()
cp.paragraph_format.space_after = Pt(8)
cr = cp.add_run("Please place the materials in one shared folder using the following organization:")
set_font(cr, 10.5)

folder_lines = [
    "01_Profile", "02_Research_Projects", "03_Publications", "04_Team",
    "05_Teaching", "06_News_Awards_Service", "07_CV", "08_Website_Ownership",
]
folder_table = doc.add_table(rows=0, cols=1)
folder_table.autofit = False
for line in folder_lines:
    c = folder_table.add_row().cells[0]
    set_cell_margins(c, top=90, bottom=90, start=180, end=180)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"Professor Website Materials / {line} /")
    set_font(r, 10, color=DARK)
set_table_geometry(folder_table, [9360])
set_table_borders(folder_table)

fh = doc.add_paragraph(style="Heading 2")
fh.add_run("Recommended filenames")
for filename in [
    "professor-headshot.jpg", "project-soil-moisture.png",
    "team-firstname-lastname.jpg", "publications.bib", "professor-cv.pdf",
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(filename)
    set_font(r, 10, color=BLUE)

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(16)
note.paragraph_format.space_after = Pt(0)
nr = note.add_run("Please do not include private, unpublished, confidential, or personally sensitive information unless it has been approved for public release.")
set_font(nr, 10.5, bold=True, color=DARK)

doc.core_properties.title = "Professor Website Information Checklist"
doc.core_properties.subject = "Academic website content collection checklist"
doc.core_properties.author = "Website Project Team"
doc.save(OUT)
print(OUT)
